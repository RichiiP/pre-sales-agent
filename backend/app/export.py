from io import BytesIO
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from .models import ExportRequest, NOTICE

BLUE, PALE, RED = "17365D", "DCE6F1", "F4CCCC"


def safe(value):
    if value is None or value == "": return "Not provided by distributor"
    text = str(value).replace("\x00", "")[:32000]
    return "'" + text if re.match(r"^[=+\-@]", text) else text


def _heading(ws, row, values):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value); cell.font = Font(color="FFFFFF", bold=True); cell.fill = PatternFill("solid", fgColor=BLUE); cell.alignment = Alignment(wrap_text=True)


def _table(ws, name, start_row, end_row, end_col):
    ref = f"A{start_row}:{ws.cell(end_row, end_col).coordinate}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table); ws.auto_filter.ref = ref; ws.freeze_panes = f"A{start_row + 1}"


def create_quote(payload: ExportRequest) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Quote Summary"
    ws["A1"] = "[COMPANY TITLE] — Product Quote"; ws["A1"].font = Font(size=18, bold=True, color=BLUE)
    meta = [("Customer name",payload.metadata.customer_name),("Opportunity / tender reference",payload.metadata.opportunity_reference),("Quote date",payload.metadata.quote_date),("Quote validity date",payload.metadata.validity_date),("Target country",payload.metadata.target_country),("Currency",payload.metadata.currency),("Prepared by",payload.metadata.prepared_by)]
    for r,(label,value) in enumerate(meta,3): ws.cell(r,1,label).font=Font(bold=True); ws.cell(r,2,safe(value))
    headers = ["Product name","Manufacturer","Model","Manufacturer part number","TD SYNNEX SKU","Unit price","Quantity","Extended price","Availability","Estimated delivery","Warranty","Confidence level"]
    header_row=12; _heading(ws,header_row,headers)
    for r,p in enumerate(payload.products,header_row+1):
        values=[p.product_name,p.manufacturer,p.model,p.manufacturer_part_number,p.sku,p.unit_price,p.requested_quantity,None,p.availability_status,p.estimated_delivery,p.warranty,p.confidence_level]
        for c,v in enumerate(values,1): ws.cell(r,c,safe(v) if c not in (6,7,8) else v)
        if p.unit_price is not None: ws.cell(r,8,f"=F{r}*G{r}")
        ws.cell(r,6).number_format=ws.cell(r,8).number_format='"$"#,##0.00'
    last=header_row+len(payload.products); subtotal=last+2
    ws.cell(subtotal,7,"Subtotal").font=Font(bold=True); ws.cell(subtotal,8,f"=SUM(H{header_row+1}:H{last})")
    ws.cell(subtotal+1,7,"Grand total").font=Font(bold=True); ws.cell(subtotal+1,8,f"=H{subtotal}"); ws.cell(subtotal+1,8).font=Font(bold=True)
    ws.cell(subtotal+3,1,NOTICE); ws.merge_cells(start_row=subtotal+3,start_column=1,end_row=subtotal+3,end_column=12)
    _table(ws,"QuoteProducts",header_row,last,12)

    tc=wb.create_sheet("Technical Compliance"); tc["A1"]="[COMPANY TITLE] — Technical Compliance"; tc["A1"].font=Font(size=18,bold=True,color=BLUE)
    th=["Customer requirement","Offered specification","Compliance status","Supporting product specification","Exception or clarification","Confidence score","Product"]
    _heading(tc,3,th); row=4
    for p in payload.products:
        specs={k.lower():v for k,v in p.technical_specifications.items()}
        for req in payload.requirements:
            offered=specs.get(req.name.lower()); exception=next((e for e in p.exceptions if e.startswith(req.name+":")),"")
            status="Requires verification" if offered is None else ("Non-compliant" if exception and req.mandatory else "Partially compliant" if exception else "Compliant")
            for c,v in enumerate([f"{req.name}: {req.value}",offered,status,offered,exception,p.confidence_score,p.product_name],1): tc.cell(row,c,safe(v))
            if status != "Compliant": tc.cell(row,3).fill=PatternFill("solid",fgColor=RED)
            row+=1
    if row == 4:
        for c,v in enumerate(["No structured requirements supplied","Not provided by distributor","Requires verification","Not provided by distributor","Add requirements before final quote",0,payload.products[0].product_name],1): tc.cell(row,c,v)
        row+=1
    _table(tc,"ComplianceItems",3,row-1,7)

    pd=wb.create_sheet("Product Details"); pd["A1"]="[COMPANY TITLE] — Product Details"; pd["A1"].font=Font(size=18,bold=True,color=BLUE)
    dh=["Product","Description","Technical specifications","Lifecycle status","Regional availability","Eligible countries","Inventory timestamp","Pricing timestamp","Data source"]
    _heading(pd,3,dh)
    for r,p in enumerate(payload.products,4):
        vals=[p.product_name,p.description,"; ".join(f"{k}: {v}" for k,v in p.technical_specifications.items()),p.lifecycle_status,p.fulfillment_region,", ".join(p.eligible_countries or []),p.inventory_timestamp,p.pricing_timestamp,p.data_source]
        for c,v in enumerate(vals,1): pd.cell(r,c,safe(v)); pd.cell(r,c).alignment=Alignment(wrap_text=True,vertical="top")
    _table(pd,"ProductDetailItems",3,3+len(payload.products),9)
    for sheet in wb:
        for col in range(1,sheet.max_column+1): sheet.column_dimensions[sheet.cell(1,col).column_letter].width=min(45,max(14,max(len(str(sheet.cell(r,col).value or "")) for r in range(1,sheet.max_row+1))+2))
    output=BytesIO(); wb.save(output); return output.getvalue()

