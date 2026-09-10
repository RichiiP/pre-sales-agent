from datetime import date, timedelta
from io import BytesIO
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app
from app.clients import MOCK_PRODUCTS
from app.matching import match_products, score_product
from app.models import Requirement, SearchRequest

client=TestClient(app)


def request(**kwargs):
    base=dict(category="Storage",requirements=[Requirement(name="capacity",value="20TB"),Requirement(name="form factor",value="3.5-inch"),Requirement(name="interface",value="SATA"),Requirement(name="speed",value="7200 RPM")],quantity=6)
    base.update(kwargs); return SearchRequest(**base)


def test_defaults_to_jamaica(): assert SearchRequest(category="Storage").country == "Jamaica"


def test_filters_lifecycle_region_and_low_scores():
    primary,regional,_=match_products(MOCK_PRODUCTS,request(allow_exceptions=True))
    assert primary and len(primary)<=5
    assert all(p.lifecycle_status == "Active" and p.orderable for p in primary)
    assert all(p.confidence_score >= 70 for p in primary)
    assert all("Jamaica" in p.eligible_countries for p in primary)
    assert all(p.product_id != "MOCK-004" for p in primary)
    assert any(p.product_id == "MOCK-021" for p in regional) is False or all("Jamaica" not in (p.eligible_countries or []) for p in regional)


def test_mandatory_failure_and_exception_policy():
    req=request(requirements=[Requirement(name="capacity",value="99TB")])
    scored=score_product(MOCK_PRODUCTS[0],req)
    assert scored.exceptions and "requested 99TB" in scored.exceptions[0]
    primary,_,_=match_products(MOCK_PRODUCTS,req)
    assert scored.product_id not in [p.product_id for p in primary]


def test_quantity_extended_price_and_missing_price():
    item=score_product(MOCK_PRODUCTS[1],request())
    assert item.quantity_shortfall == 2 and item.extended_price == item.unit_price*6
    missing=score_product(MOCK_PRODUCTS[18],request(requirements=[]))
    assert missing.unit_price is None and missing.extended_price is None


def test_api_and_export_have_no_credentials(monkeypatch):
    monkeypatch.setenv("TD_SYNNEX_CLIENT_SECRET","super-secret-value")
    result=client.post("/api/search",json=request().model_dump(mode="json")); assert result.status_code == 200
    assert "super-secret-value" not in result.text
    product=result.json()["recommendations"][0]
    payload={"metadata":{"customer_name":"Acme","opportunity_reference":"OP-1","prepared_by":"Engineer","quote_date":str(date.today()),"validity_date":str(date.today()+timedelta(days=14)),"target_country":"Jamaica","currency":"USD"},"products":[product],"requirements":request().model_dump(mode="json")["requirements"]}
    exported=client.post("/api/export",json=payload); assert exported.status_code == 200
    assert b"super-secret-value" not in exported.content
    wb=load_workbook(BytesIO(exported.content),data_only=False)
    assert wb.sheetnames == ["Quote Summary","Technical Compliance","Product Details"]
    assert wb["Quote Summary"]["H13"].value == "=F13*G13"
    assert wb["Quote Summary"]["H15"].value == "=SUM(H13:H13)"
    assert wb["Quote Summary"]["E13"].value == product["sku"]
